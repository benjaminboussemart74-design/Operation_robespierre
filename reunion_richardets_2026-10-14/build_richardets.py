"""Présentation de la réunion publique du quartier des Richardets (14 octobre 2026).

La présentation reprend la forme graphique des réunions précédentes : elle part du
fichier du Mont d'Est (modele/PPT_Mont_dEst_V3.pptx), en conserve les diapositives
utiles comme gabarits (logo, QR code, couleurs, ombres des titres), remplace les
textes et la vue aérienne, puis supprime le reste.

Fond de la diapositive de titre : orthophotographie IGN (Géoplateforme, WMTS).

Usage : python3 build_richardets.py
"""
import copy
import io
import math
import os
import subprocess
import time
from concurrent.futures import ThreadPoolExecutor

from PIL import Image, ImageEnhance
from pptx import Presentation

ICI = os.path.dirname(os.path.abspath(__file__))
MODELE = os.path.join(ICI, "modele", "PPT_Mont_dEst_V3.pptx")
SORTIE = os.path.join(ICI, "PPT_Richardets_V0.pptx")
CACHE = os.path.join(ICI, ".cache")

QUARTIER = "QUARTIER DES RICHARDETS"
DATE_LIEU = "Mercredi 14 octobre 2026 — 19 h — Victoria Park"

# Numéros (base 1) des diapositives du Mont d'Est servant de gabarits
G_QR_ACCUEIL, G_TITRE, G_SOMMAIRE, G_QUESTIONS, G_QR_SORTIE = 1, 2, 3, 28, 29

# Sommaire provisoire, à remplacer par le déroulé définitif
SOMMAIRE = [
    "Proximité et démocratie locale",
    "Les écoles et le périscolaire",
    "Les Richardets : services publics de proximité",
    "Propreté dans le quartier des Richardets",
    "Les Richardets : sécurité et tranquillité résidentielle",
    "Le logement social",
    "Les transports : les bus",
]

# Centre de la vue aérienne (résidence des Richardets, géocodage IGN) et largeur au sol
CENTRE_LON, CENTRE_LAT, LARGEUR_SOL_M = 2.5635, 48.8318, 900


# ---------------------------------------------------------------------------
# Orthophotographie IGN
# ---------------------------------------------------------------------------
WMTS = ("https://data.geopf.fr/wmts?SERVICE=WMTS&REQUEST=GetTile&VERSION=1.0.0"
        "&LAYER=ORTHOIMAGERY.ORTHOPHOTOS&STYLE=normal&TILEMATRIXSET=PM_0_19"
        "&TILEMATRIX={z}&TILEROW={y}&TILECOL={x}&FORMAT=image/jpeg")
R = 6378137.0


def merc(lon, lat):
    return (math.radians(lon) * R, math.log(math.tan(math.pi / 4 + math.radians(lat) / 2)) * R)


def tuile(z, x, y):
    f = os.path.join(CACHE, f"{z}_{x}_{y}.jpg")
    for essai in range(8):
        if os.path.exists(f):
            return f
        r = subprocess.run(["curl", "-sS", "-m", "60", WMTS.format(z=z, x=x, y=y)], capture_output=True)
        if r.returncode == 0 and r.stdout[:3] == b"\xff\xd8\xff":   # vraie image JPEG
            open(f + ".tmp", "wb").write(r.stdout)
            os.replace(f + ".tmp", f)
        else:
            time.sleep(1 + essai)
    raise RuntimeError(f"Tuile {z}/{x}/{y} indisponible")


def vue_aerienne(lon, lat, largeur_m, taille_px):
    """Assemble les tuiles IGN et recadre sur l'emprise demandée, au format de taille_px."""
    largeur, hauteur = taille_px
    cx, cy = merc(lon, lat)
    w = largeur_m / math.cos(math.radians(lat))
    h = w * hauteur / largeur
    x0, y1 = cx - w / 2, cy + h / 2
    z = min(19, math.ceil(math.log2(2 * math.pi * R / 256 / (w / largeur))))
    taille = 2 * math.pi * R / 2 ** z
    tx0, tx1 = int((x0 + math.pi * R) // taille), int((x0 + w + math.pi * R) // taille)
    ty0, ty1 = int((math.pi * R - y1) // taille), int((math.pi * R - (y1 - h)) // taille)
    os.makedirs(CACHE, exist_ok=True)
    tuiles = [(x, y) for x in range(tx0, tx1 + 1) for y in range(ty0, ty1 + 1)]
    mosaique = Image.new("RGB", ((tx1 - tx0 + 1) * 256, (ty1 - ty0 + 1) * 256))
    with ThreadPoolExecutor(max_workers=12) as pool:
        for (x, y), f in zip(tuiles, pool.map(lambda t: tuile(z, *t), tuiles)):
            mosaique.paste(Image.open(f).convert("RGB"), ((x - tx0) * 256, (y - ty0) * 256))
    e = 256 / taille
    ox, oy = tx0 * taille - math.pi * R, math.pi * R - ty0 * taille
    boite = ((x0 - ox) * e, (oy - y1) * e, (x0 + w - ox) * e, (oy - y1 + h) * e)
    img = mosaique.crop(tuple(round(v) for v in boite)).resize((largeur, hauteur), Image.LANCZOS)
    return ImageEnhance.Color(img).enhance(0.85)


# ---------------------------------------------------------------------------
# Manipulation des diapositives
# ---------------------------------------------------------------------------
def forme(diapo, nom):
    for sh in diapo.shapes:
        if sh.name == nom:
            return sh
    raise KeyError(nom)


def remplacer_texte(diapo, nom, *paragraphes):
    """Remplace le texte d'une forme en conservant la mise en forme de son premier segment."""
    tf = forme(diapo, nom).text_frame
    modele_p = tf.paragraphs[0]._p
    modele_r = copy.deepcopy(modele_p.r_lst[0])
    for p in list(tf._txBody.p_lst):
        tf._txBody.remove(p)
    for texte in paragraphes:
        p = copy.deepcopy(modele_p)
        for r in p.r_lst:
            p.remove(r)
        r = copy.deepcopy(modele_r)
        r.t.text = texte
        p.insert(1 if p.pPr is not None else 0, r)
        tf._txBody.append(p)


def notes(diapo, texte):
    diapo.notes_slide.notes_text_frame.text = texte


def remplacer_image(diapo, nom, image):
    pic = forme(diapo, nom)
    rid = pic._element.blipFill.blip.rEmbed
    tampon = io.BytesIO()
    image.save(tampon, "JPEG", quality=88)
    diapo.part.related_part(rid)._blob = tampon.getvalue()


def ordonner(prs, numeros):
    """Ne garde que les diapositives indiquées (numéros base 1), dans cet ordre."""
    liste = prs.slides._sldIdLst
    ids = list(liste)
    for sld in ids:
        liste.remove(sld)
    for n in numeros:
        liste.append(ids[n - 1])
    for sld in ids:
        if sld not in liste:
            prs.part.drop_rel(sld.rId)


def dupliquer(prs, numero):
    """Copie la diapositive `numero` (base 1) à la fin de la présentation, images comprises."""
    src = prs.slides[numero - 1]
    dst = prs.slides.add_slide(src.slide_layout)
    for sh in list(dst.shapes):
        sh._element.getparent().remove(sh._element)
    correspondance = {}
    for rid, rel in src.part.rels.items():
        if "image" in rel.reltype:
            correspondance[rid] = dst.part.relate_to(rel.target_part, rel.reltype)
    for el in src.shapes._spTree.iterchildren():
        if el.tag.endswith(("}sp", "}pic", "}grpSp", "}cxnSp", "}graphicFrame")):
            nouveau = copy.deepcopy(el)
            for blip in nouveau.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}blip"):
                ancien = blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
                if ancien in correspondance:
                    blip.set("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed",
                             correspondance[ancien])
            dst.shapes._spTree.append(nouveau)
    return dst


# ---------------------------------------------------------------------------
# Construction
# ---------------------------------------------------------------------------
def construire():
    prs = Presentation(MODELE)
    s = {n: prs.slides[n - 1] for n in range(1, len(prs.slides) + 1)}

    # Accueil et sortie : QR code et fiche de renseignement
    for n, texte in ((G_QR_ACCUEIL, "Diapositive affichée pendant l'accueil : inviter les habitants à "
                                    "flasher le QR code ou à remplir la fiche de renseignement papier."),
                     (G_QR_SORTIE, "Diapositive laissée à l'écran à la sortie : rappeler le QR code et "
                                   "la fiche de renseignement papier.")):
        remplacer_texte(s[n], "Text 0", "RÉUNION PUBLIQUE · " + QUARTIER)
        notes(s[n], texte)

    # Titre
    titre = s[G_TITRE]
    remplacer_texte(titre, "Text 4", "QUARTIER DES", "RICHARDETS")
    remplacer_texte(titre, "Text 5", DATE_LIEU)
    remplacer_image(titre, "Image 0", vue_aerienne(CENTRE_LON, CENTRE_LAT, LARGEUR_SOL_M, (1920, 1080)))
    notes(titre, "Accueil des habitants. Intervenants : Vincent Monnier, Maire de Noisy-le-Grand ; "
                 "élu référent du quartier : à compléter.")

    # Sommaire : on garde autant de lignes que de parties
    sommaire = s[G_SOMMAIRE]
    for i in range(1, 9):
        if i <= len(SOMMAIRE):
            remplacer_texte(sommaire, f"Partie {i}", SOMMAIRE[i - 1])
        else:
            for nom in (f"Numero {i}", f"Partie {i}", f"Separateur {i - 1}"):
                el = forme(sommaire, nom)._element
                el.getparent().remove(el)
    notes(sommaire, f"Présenter les {len(SOMMAIRE)} parties de la soirée. Sommaire provisoire, "
                    "à aligner sur le déroulé définitif.")

    notes(s[G_QUESTIONS], "Temps d'échange avec la salle.")

    ordonner(prs, [G_QR_ACCUEIL, G_TITRE, G_SOMMAIRE, G_QUESTIONS, G_QR_SORTIE])
    prs.save(SORTIE)
    print("Écrit :", SORTIE)


if __name__ == "__main__":
    construire()
