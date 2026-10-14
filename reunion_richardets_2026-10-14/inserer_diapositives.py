"""Insère des diapositives d'une présentation générée dans une présentation retouchée à la main,
sans modifier les autres diapositives de cette dernière.

Usage : python3 inserer_diapositives.py CIBLE.pptx SOURCE.pptx SORTIE.pptx POSITION N1 [N2 ...]
  POSITION : numéro (base 1) de la diapositive de CIBLE après laquelle insérer ;
  N1, N2… : numéros (base 1) des diapositives de SOURCE à copier, dans cet ordre.
"""
import copy
import io
import sys

from pptx import Presentation

R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def copier(cible, src):
    dst = cible.slides.add_slide(cible.slides[0].slide_layout)
    for ph in list(dst.placeholders):
        ph._element.getparent().remove(ph._element)
    # images : on recrée chaque relation dans la cible
    corresp = {}
    for rid, rel in src.part.rels.items():
        if "image" in rel.reltype:
            img = dst.part.package.get_or_add_image_part(io.BytesIO(rel.target_part.blob))
            corresp[rid] = dst.part.relate_to(img, rel.reltype)
    for el in src.shapes._spTree.iterchildren():
        if el.tag.endswith(("}sp", "}pic", "}grpSp", "}cxnSp", "}graphicFrame")):
            n = copy.deepcopy(el)
            for blip in n.iter(A + "blip"):
                if blip.get(R + "embed") in corresp:
                    blip.set(R + "embed", corresp[blip.get(R + "embed")])
            dst.shapes._spTree.append(n)
    # fond de la diapositive
    bg = src._element.cSld.bg
    if bg is not None:
        dst._element.cSld.insert(0, copy.deepcopy(bg))
    # notes de l'orateur
    if src.has_notes_slide and src.notes_slide.notes_text_frame is not None:
        ns = dst.notes_slide
        if ns.notes_text_frame is None:
            modele = next(s for s in cible.slides if s.has_notes_slide and s.notes_slide.notes_placeholder)
            ns.shapes._spTree.append(copy.deepcopy(modele.notes_slide.notes_placeholder._element))
        ns.notes_text_frame.text = src.notes_slide.notes_text_frame.text
    return dst


def main(cible_f, source_f, sortie_f, position, numeros):
    cible, source = Presentation(cible_f), Presentation(source_f)
    for n in numeros:
        copier(cible, source.slides[n - 1])
    liste = cible.slides._sldIdLst
    nouvelles = list(liste)[-len(numeros):]
    for i, sld in enumerate(nouvelles):
        liste.remove(sld)
        liste.insert(position + i, sld)
    cible.save(sortie_f)
    print("Écrit :", sortie_f, "-", len(liste), "diapositives")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2], sys.argv[3], int(sys.argv[4]), [int(x) for x in sys.argv[5:]])
